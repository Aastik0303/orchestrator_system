from __future__ import annotations

import hashlib
import logging
import re
import time
from pathlib import Path
from threading import Lock
from typing import Any
from uuid import uuid4

from docx import Document as DocxDocument
from langchain_core.documents import Document as LangChainDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

from app.config import get_settings
from app.guardrails.detectors import RETRIEVAL_DETECTORS, scan
from app.observability.telemetry import maybe_span
from app.services.runtime_store import RuntimeStore, runtime_store

logger = logging.getLogger("orchestrator.rag")
BACKEND_RETRY_SECONDS = 60


class EmbeddingUnavailable(RuntimeError):
    """The embedding backend cannot be loaded (missing package or model)."""


class EmbeddingService:
    def __init__(self) -> None:
        self._backend: Any | None = None
        self._cache_key: tuple[str, str, str, bool, int] | None = None
        self._lock = Lock()
        self._failed_at: float | None = None
        self._failure: str | None = None

    @property
    def model_identifier(self) -> str:
        settings = get_settings()
        return f"huggingface:{settings.embedding_model}"

    def embed(self, text: str) -> list[float]:
        return self.embed_query(text)

    def embed_query(self, text: str) -> list[float]:
        with maybe_span("retrieval", "embed_query"):
            vector = self._get_backend().embed_query(text)
        return self._validate_vector(vector)

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        if not texts:
            return []
        with maybe_span("retrieval", "embed_documents", count=len(texts)):
            vectors = self._get_backend().embed_documents(texts)
        return [self._validate_vector(vector) for vector in vectors]

    def _get_backend(self):
        settings = get_settings()
        provider = settings.embedding_provider.strip().lower()
        if provider not in {"huggingface", "sentence_transformers"}:
            raise ValueError(
                "EMBEDDING_PROVIDER must be 'huggingface'; the hash embedding adapter is no longer supported."
            )
        cache_key = (
            provider,
            settings.embedding_model,
            settings.embedding_device,
            settings.embedding_local_files_only,
            settings.embedding_batch_size,
        )
        if self._backend is not None and self._cache_key == cache_key:
            return self._backend
        # Do not retry a failed backend load on every call (a missing package
        # would otherwise cost an import attempt per request).
        if self._failed_at is not None and time.monotonic() - self._failed_at < BACKEND_RETRY_SECONDS:
            raise EmbeddingUnavailable(self._failure or "Embedding backend unavailable.")
        with self._lock:
            if self._backend is None or self._cache_key != cache_key:
                try:
                    from langchain_huggingface import HuggingFaceEmbeddings
                except ImportError as exc:
                    self._failed_at = time.monotonic()
                    self._failure = f"Embedding backend is not installed ({exc.name})."
                    logger.warning("embedding_backend_unavailable")
                    raise EmbeddingUnavailable(self._failure) from None

                if settings.embedding_torch_threads > 0:
                    # Many small concurrent queries: one intra-op thread per
                    # call avoids oversubscription (measured: 8.0 vs 13.7 ms
                    # per call at 10 concurrent callers, 1 vs 16 threads).
                    import torch

                    torch.set_num_threads(settings.embedding_torch_threads)
                self._backend = HuggingFaceEmbeddings(
                    model=settings.embedding_model,
                    model_kwargs={
                        "device": settings.embedding_device,
                        "local_files_only": settings.embedding_local_files_only,
                    },
                    encode_kwargs={
                        "normalize_embeddings": True,
                        "batch_size": settings.embedding_batch_size,
                    },
                    query_encode_kwargs={"normalize_embeddings": True},
                    show_progress=False,
                )
                self._cache_key = cache_key
                self._failed_at = None
        return self._backend

    @staticmethod
    def _validate_vector(vector: Any) -> list[float]:
        normalized = [float(value) for value in vector]
        expected = get_settings().embedding_dimensions
        if len(normalized) != expected:
            raise ValueError(
                f"Embedding model returned {len(normalized)} dimensions; expected {expected}."
            )
        return normalized


embedding_service = EmbeddingService()


def extract_document(path: Path) -> list[tuple[int | None, str]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(path)
        return [
            (page_number, (page.extract_text() or "").strip())
            for page_number, page in enumerate(reader.pages, start=1)
            if (page.extract_text() or "").strip()
        ]
    if suffix == ".docx":
        document = DocxDocument(path)
        text = "\n".join(
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        )
        return [(None, text)] if text else []
    if suffix in {".txt", ".md", ".csv"}:
        text = path.read_text(encoding="utf-8", errors="ignore").strip()
        return [(None, text)] if text else []
    if suffix in {".xlsx", ".xls"}:  # .xls needs xlrd
        import pandas as pd

        workbook = pd.read_excel(path, sheet_name=None)
        sections = [f"Sheet: {name}\n{frame.to_csv(index=False)}" for name, frame in workbook.items()]
        return [(None, "\n\n".join(sections))]
    return []


def chunk_pages(
    pages: list[tuple[int | None, str]], *, chunk_size: int, overlap: int
) -> list[dict[str, Any]]:
    if chunk_size <= 0 or overlap < 0 or overlap >= chunk_size:
        raise ValueError("Chunk size must be positive and overlap must be smaller than chunk size.")
    documents: list[LangChainDocument] = []
    for page_number, raw_text in pages:
        text = re.sub(r"[ \t]+", " ", raw_text)
        text = re.sub(r"\n{3,}", "\n\n", text).strip()
        if text:
            documents.append(
                LangChainDocument(
                    page_content=text,
                    metadata={"page_number": page_number},
                )
            )
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
        add_start_index=True,
        strip_whitespace=True,
    )
    split_documents = splitter.split_documents(documents)
    return [
        {
            "id": f"chunk_{uuid4().hex[:16]}",
            "chunk_index": index,
            "page_number": document.metadata.get("page_number"),
            "content": document.page_content,
        }
        for index, document in enumerate(split_documents)
        if document.page_content.strip()
    ]


def index_document(
    document_id: str,
    *,
    user_id: str = "local-user",
    store: RuntimeStore = runtime_store,
) -> dict[str, Any]:
    document = store.get_document(document_id, user_id=user_id)
    if not document:
        raise FileNotFoundError(f"Document {document_id} was not found.")
    path = Path(document["storage_path"])
    if not path.is_file():
        store.update_document(document_id, status="failed", error="Stored file is missing.")
        raise FileNotFoundError("Stored document file is missing.")

    settings = get_settings()
    try:
        store.update_document(document_id, status="extracting", error=None)
        pages = extract_document(path)
        if not pages:
            raise ValueError("No readable text could be extracted from the document.")
        store.update_document(document_id, status="chunking")
        chunks = chunk_pages(
            pages,
            chunk_size=settings.rag_chunk_size,
            overlap=settings.rag_chunk_overlap,
        )
        if not chunks:
            raise ValueError("Document extraction produced no indexable chunks.")
        for chunk in chunks:
            # Stable ids: re-indexing a document keeps its chunk ids (and so
            # its citations) unchanged.
            chunk["id"] = "chunk_" + hashlib.sha256(f"{document_id}:{chunk['chunk_index']}".encode()).hexdigest()[:16]
        store.update_document(document_id, status="embedding")
        vectors = embedding_service.embed_documents([chunk["content"] for chunk in chunks])
        for chunk, vector in zip(chunks, vectors, strict=True):
            chunk["embedding"] = vector
        store.replace_chunks(document_id, chunks)
        store.update_document(
            document_id,
            status="indexed",
            chunk_count=len(chunks),
            embedding_model=embedding_service.model_identifier,
            error=None,
        )
    except Exception as exc:
        store.update_document(document_id, status="failed", error=str(exc)[:500])
        raise
    return store.get_document(document_id, user_id=user_id) or document


def search_knowledge(
    query: str,
    *,
    user_id: str = "local-user",
    project_id: str | None = "default",
    top_k: int | None = None,
    threshold: float | None = None,
    document_ids: list[str] | None = None,
    store: RuntimeStore = runtime_store,
) -> list[dict[str, Any]]:
    """Vector search over the caller's current-model documents.

    Documents indexed with a different embedding model are skipped (and
    reported via `stale_documents`) rather than re-embedded inside the query
    path; re-indexing belongs to the explicit /reindex endpoint or a worker.

    When `document_ids` scopes the search to documents the user attached, the
    best chunks are always returned: the user is asking about those files, and
    broad questions ("summarize my resume", "what is his education") score far
    below the global similarity threshold against every chunk.
    """
    settings = get_settings()
    current_model = embedding_service.model_identifier
    documents = store.list_documents(user_id=user_id, project_id=project_id)
    allowed_ids = set(document_ids or [])
    current_documents = [
        document
        for document in documents
        if document["status"] == "indexed"
        and document.get("embedding_model") == current_model
        and (not document_ids or document["id"] in allowed_ids)
    ]
    if not current_documents:
        return []
    scoped = bool(document_ids)
    if threshold is None:
        threshold = 0.0 if scoped else settings.rag_similarity_threshold

    with maybe_span("retrieval", "vector_search", documents=len(current_documents)):
        results = store.search_chunks(
            embedding_service.embed_query(query),
            user_id=user_id,
            project_id=project_id,
            top_k=top_k or settings.rag_top_k,
            threshold=threshold,
            document_ids=[document["id"] for document in current_documents],
            embedding_model=current_model,
        )
    if results and settings.rag_relative_score_cutoff > 0 and not scoped:
        best = results[0]["similarity"]
        results = [item for item in results if item["similarity"] >= best * settings.rag_relative_score_cutoff]
    for result in results:
        result["prompt_injection_detected"] = bool(scan(result["content"], RETRIEVAL_DETECTORS))
    return results


def stale_documents(*, user_id: str, project_id: str | None, store: RuntimeStore = runtime_store) -> list[str]:
    current_model = embedding_service.model_identifier
    return [
        document["id"]
        for document in store.list_documents(user_id=user_id, project_id=project_id)
        if document["status"] == "indexed" and document.get("embedding_model") != current_model
    ]
